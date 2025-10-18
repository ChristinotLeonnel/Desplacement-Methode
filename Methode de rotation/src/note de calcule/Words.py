"""
Note de calcul : Génération de documents Word avec Python

Ce module fournit des fonctions pour générer automatiquement des notes de calcul ou des rapports techniques au format Microsoft Word (.docx) à partir de données structurées (par exemple, issues de calculs de structure, d'analyses ou de résultats d'essais).

Bibliothèque utilisée :
    - python-docx : Permet de créer et de manipuler des documents Word.

Fonctionnalités principales :
    - Création d'un document Word avec titre, sections, paragraphes et tableaux.
    - Insertion de résultats de calculs, de textes explicatifs et de tableaux de données.
    - Mise en forme automatique (titres, styles, alignements).
    - Export du document finalisé.

Exemple d'utilisation :
    from Words import generer_note_de_calcul

    donnees = {
        "titre": "Note de calcul - Poutre BA",
        "sections": [
            {
                "titre": "1. Introduction",
                "contenu": "Cette note présente le calcul de la poutre BA selon l'Eurocode 2."
            },
            {
                "titre": "2. Hypothèses",
                "contenu": "Les hypothèses de calcul sont les suivantes :\n- Béton C25/30\n- Acier Fe500"
            },
            {
                "titre": "3. Résultats",
                "tableau": {
                    "en_tete": ["Section", "Effort normal (kN)", "Moment (kNm)", "Flèche (mm)"],
                    "lignes": [
                        ["A-A", 0, 45.2, 3.1],
                        ["B-B", 0, 32.8, 2.4]
                    ]
                }
            }
        ]
    }
    generer_note_de_calcul(donnees, "Note_de_calcul_Poutre_BA.docx")

"""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generer_note_de_calcul(donnees, chemin_fichier):
    """
    Génère une note de calcul Word à partir d'un dictionnaire de données.

    Args:
        donnees (dict): Dictionnaire contenant le titre, les sections, et éventuellement des tableaux.
        chemin_fichier (str): Chemin de sauvegarde du fichier Word généré.
    """
    doc = Document()

    # Titre principal
    titre = donnees.get("titre", "Note de calcul")
    titre_paragraph = doc.add_heading(titre, level=0)
    titre_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()  # Saut de ligne

    # Parcours des sections
    for section in donnees.get("sections", []):
        # Titre de section
        titre_section = section.get("titre", "")
        if titre_section:
            doc.add_heading(titre_section, level=1)

        # Contenu texte
        contenu = section.get("contenu", "")
        if contenu:
            p = doc.add_paragraph(contenu)
            p.style.font.size = Pt(11)

        # Tableau éventuel
        tableau = section.get("tableau")
        if tableau:
            en_tete = tableau.get("en_tete", [])
            lignes = tableau.get("lignes", [])
            table = doc.add_table(rows=1, cols=len(en_tete))
            table.style = "Table Grid"
            hdr_cells = table.rows[0].cells
            for i, titre_col in enumerate(en_tete):
                hdr_cells[i].text = str(titre_col)
            for ligne in lignes:
                row_cells = table.add_row().cells
                for i, val in enumerate(ligne):
                    row_cells[i].text = str(val)
            doc.add_paragraph()  # Saut de ligne après le tableau

    # Pied de page (optionnel)
    section = doc.sections[0]
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.text = "Généré automatiquement - Méthode de Rotation"
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Sauvegarde du document
    doc.save(chemin_fichier)
